"""Deterministic text extraction for supported RFQ documents; no OCR or AI inference."""
from __future__ import annotations
import io
import re
import zipfile
from pathlib import Path
from typing import Any
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_EXTRACTED_CHARS = 2_000_000


def extract_document(filename: str, content: bytes) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    extension = Path(filename).suffix.lower()
    segments: list[dict[str, Any]] = []
    warnings: list[str] = []
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported. Upload an unlocked PDF.")
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            segments.append({"source": f"Page {index}", "text": text})
        if not any(s["text"].strip() for s in segments):
            raise ValueError("No selectable text was found. Scanned PDFs require OCR, which is not enabled in this release.")
    elif extension == ".docx":
        _validate_office_archive(content)
        doc = WordDocument(io.BytesIO(content))
        for index, paragraph in enumerate(doc.paragraphs, start=1):
            if paragraph.text.strip():
                segments.append({"source": f"Paragraph {index}", "text": paragraph.text.strip()})
        for table_index, table in enumerate(doc.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                text = _label_value_row(values)
                if text.strip(" |"):
                    segments.append({"source": f"Table {table_index}, row {row_index}", "text": text})
    elif extension == ".xlsx":
        _validate_office_archive(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    populated = [(cell.coordinate, str(cell.value).strip()) for cell in row if cell.value is not None and str(cell.value).strip()]
                    if populated:
                        normalized = _label_value_row([value for _, value in populated])
                        details = "; ".join(f"{coordinate}={value}" for coordinate, value in populated)
                        segments.append({"source": f"Sheet {sheet.title}, row {row[0].row} ({details})", "text": normalized})
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or XLSX document.")

    text = "\n".join(f"[{segment['source']}] {segment['text']}" for segment in segments)
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS]
        warnings.append("Extracted text was truncated at the configured character limit.")
    if not text.strip():
        raise ValueError("No readable text or cell values were found in this document.")
    fields = detect_fields(segments)
    if not fields:
        warnings.append("No supported labeled requirements were detected. You can still review the source text; no values were invented.")
    return text, segments, fields, warnings


def detect_fields(segments: list[dict[str, str]]) -> list[dict[str, Any]]:
    patterns = [
        ("part_number", "Part number", re.compile(r"\b(?:part\s*(?:number|no\.?|#)|part\s*id|item\s*code|drawing\s*(?:number|no\.?))\s*[:=#-]\s*([^;|,\n]+)", re.I)),
        ("material", "Material", re.compile(r"\bmaterial\s*[:=-]\s*([^;|,\n]+)", re.I)),
        ("dimensions", "Dimensions", re.compile(r"\b(?:overall\s+)?dimensions?\s*[:=-]\s*([^;|\n]+)", re.I)),
        ("tolerance", "Tolerance", re.compile(r"\b(?:general\s+)?tolerances?\s*[:=-]\s*([^;|,\n]+)", re.I)),
        ("quantity", "Quantity", re.compile(r"\b(?:quantity|qty)\s*[:=-]\s*([0-9]+(?:\.[0-9]+)?(?:\s*[a-zA-Z]+)?)", re.I)),
        ("finish", "Surface finish", re.compile(r"\b(?:surface\s+)?finish\s*[:=-]\s*([^;|,\n]+)", re.I)),
        ("delivery_date", "Delivery requirement", re.compile(r"\b(?:delivery\s+date|required\s+by|delivery)\s*[:=-]\s*([^;|\n]+)", re.I)),
    ]
    found: dict[str, dict[str, Any]] = {}
    for segment in segments:
        for key, label, pattern in patterns:
            if key in found:
                continue
            match = pattern.search(segment["text"])
            if match:
                value = match.group(1).strip().strip(" .")
                # Drawing title blocks and dimension callouts often contain words that
                # resemble labels. Reject obvious metadata and implausibly weak matches.
                if not _plausible_candidate(key, value):
                    continue
                found[key] = {"key": key, "label": label, "value": value, "source": segment["source"], "confidence": 0.55, "review_status": "needs_review"}
    return list(found.values())


def _validate_office_archive(content: bytes) -> None:
    """Bound ZIP expansion and entry count before parsing Office XML."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000:
                raise ValueError("Office documents with more than 5,000 archive entries are not supported.")
            if sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise ValueError("The expanded Office document exceeds the 50 MB processing limit.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError("Encrypted Office documents are not supported.")
    except zipfile.BadZipFile:
        raise ValueError("The Office document is not a valid DOCX/XLSX archive.") from None


def _label_value_row(values: list[str]) -> str:
    """Turn adjacent spreadsheet/table label-value cells into extractable label: value pairs."""
    cleaned = [value.strip() for value in values if value.strip()]
    if len(cleaned) >= 2 and len(cleaned) % 2 == 0:
        pairs = []
        for index in range(0, len(cleaned), 2):
            label, value = cleaned[index:index + 2]
            if re.fullmatch(r"[A-Za-z][A-Za-z0-9 _/#().-]{1,60}", label) and len(value) <= 2000:
                pairs.append(f"{label}: {value}")
            else:
                pairs.append(f"{label} | {value}")
        return " | ".join(pairs)
    return " | ".join(cleaned)


def _plausible_candidate(key: str, value: str) -> bool:
    value = value.strip()
    if len(value) < 2 or len(value) > 200:
        return False
    normalized = re.sub(r"[^a-z]", "", value.lower())
    title_block_noise = {"approvedby", "checkedby", "drawnby", "date", "scale", "sheet", "revision", "rev", "approved", "signature", "name"}
    if normalized in title_block_noise or value.endswith(":"):
        return False
    if key == "material":
        # Avoid treating drawing sheet counters and dimension values as material specs.
        if re.fullmatch(r"(?:\d+\s+of\s+\d+|[\d.]+|[A-Z]\s*[A-Z])", value, re.I):
            return False
        if not re.search(r"[A-Za-z]", value):
            return False
    if key == "finish" and normalized in {"approvedby", "checkedby", "drawnby", "notapplicable"}:
        return False
    if key == "part_number" and re.fullmatch(r"[A-Z]\s+[A-Z]", value, re.I):
        return False
    return True

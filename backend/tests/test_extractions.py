from io import BytesIO

from docx import Document
from openpyxl import Workbook

from test_auth_customers import make_client, register


def docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Requirements"
    sheet.append(["Part Number", "PN-77"])
    sheet.append(["Material", "Aluminium 6061"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_rfq(client, headers, reference: str = "RFQ-EXT-01") -> str:
    customer_response = client.post(
        "/customers",
        headers=headers,
        json={"name": "Extraction Test Customer"},
    )
    assert customer_response.status_code == 201, customer_response.text
    customer_id = customer_response.json()["id"]

    rfq_response = client.post(
        "/rfqs",
        headers=headers,
        json={
            "customer_id": customer_id,
            "reference": reference,
            "title": "Extraction test RFQ",
            "quantity": 25,
        },
    )
    assert rfq_response.status_code == 201, rfq_response.text
    return rfq_response.json()["id"]


def test_docx_extraction_review_and_org_isolation():
    client, engine = make_client()
    try:
        auth = register(client, "Extraction Org", "extract@example.com")
        headers = {"Authorization": f"Bearer {auth['access_token']}"}
        rfq_id = create_rfq(client, headers)

        response = client.post(
            "/extractions",
            headers=headers,
            data={"rfq_id": rfq_id},
            files={
                "file": (
                    "rfq.docx",
                    docx_bytes(
                        "Part Number: BR-42\n"
                        "Material: Aluminium 6061\n"
                        "Quantity: 25"
                    ),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 201, response.text
        doc = response.json()
        assert doc["status"] == "needs_review"
        assert doc["file_type"] == "docx"
        assert {field["key"] for field in doc["fields"]} >= {
            "part_number",
            "material",
            "quantity",
        }
        assert all(field["source"] for field in doc["fields"])

        approved = [
            {
                "key": field["key"],
                "value": field["value"],
                "review_status": "approved",
            }
            for field in doc["fields"]
        ]
        reviewed = client.patch(
            f"/extractions/{doc['id']}/review",
            headers=headers,
            json={"fields": approved},
        )
        assert reviewed.status_code == 200, reviewed.text
        assert reviewed.json()["status"] == "reviewed"

        other = register(client, "Other Extraction Org", "otherextract@example.com")
        other_headers = {"Authorization": f"Bearer {other['access_token']}"}
        assert client.get("/extractions", headers=other_headers).json() == []
        assert (
            client.get(f"/extractions/{doc['id']}", headers=other_headers).status_code
            == 404
        )
    finally:
        client.close()
        engine.dispose()


def test_xlsx_extraction_and_upload_validation():
    client, engine = make_client()
    try:
        auth = register(client, "Spreadsheet Org", "sheet@example.com")
        headers = {"Authorization": f"Bearer {auth['access_token']}"}
        rfq_id = create_rfq(client, headers, reference="RFQ-XLSX-01")

        response = client.post(
            "/extractions",
            headers=headers,
            data={"rfq_id": rfq_id},
            files={
                "file": (
                    "spec.xlsx",
                    xlsx_bytes(),
                    "application/octet-stream",
                )
            },
        )
        assert response.status_code == 201, response.text
        assert "Sheet Requirements" in response.json()["extracted_text"]
        fields = {
            field["key"]: field["value"]
            for field in response.json()["fields"]
        }
        assert fields["part_number"] == "PN-77"
        assert fields["material"] == "Aluminium 6061"

        # Include the required RFQ ID so these requests reach file validation
        # instead of failing first with a missing-field 422 response.
        unsupported = client.post(
            "/extractions",
            headers=headers,
            data={"rfq_id": rfq_id},
            files={
                "file": (
                    "notes.txt",
                    b"Material: steel",
                    "text/plain",
                )
            },
        )
        assert unsupported.status_code == 415, unsupported.text

        invalid = client.post(
            "/extractions",
            headers=headers,
            data={"rfq_id": rfq_id},
            files={
                "file": (
                    "fake.pdf",
                    b"not a pdf",
                    "application/pdf",
                )
            },
        )
        assert invalid.status_code == 415, invalid.text
    finally:
        client.close()
        engine.dispose()
